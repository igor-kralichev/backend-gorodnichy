from __future__ import annotations

import io
import json
import zipfile
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from uuid import UUID
from xml.etree.ElementTree import ParseError

import orjson
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.security import ActorContext
from app.modules.access.application.service import AuthorizationService
from app.modules.entities.infrastructure.models import (
    EntityFieldModel,
    EntitySchemaModel,
    EntitySchemaVersionModel,
)
from app.modules.change_sets.api.schemas import ChangeSetCreate, ChangeSetItemCreate
from app.modules.change_sets.application.service import (
    ChangeSetConflict,
    ChangeSetService,
    ChangeSetValidationError,
)
from app.modules.excel.api.schemas import (
    ExcelCellIssue,
    ExcelMapping,
    ExcelPreviewRead,
    ExcelPreviewRow,
    ExcelImportPlanDecision,
    ExcelImportPlanRead,
    ExcelExportRequest,
    ImportProfileCreate,
    ImportProfileRead,
)
from app.modules.objects.api.schemas import (
    EntityObjectCreate,
    GeoJsonGeometry,
    RegistryTreeSearch,
)
from app.modules.objects.application.service import (
    RuntimeEntityNotFound,
    RuntimeObjectService,
    RuntimeValidationError,
)
from app.shared.db.models import (
    ChangeSetItemModel,
    ChangeSetModel,
    EntityObjectModel,
    ExcelImportPlanModel,
    ImportProfileModel,
)


class ExcelValidationError(Exception):
    """XLSX не соответствует ограничениям или выбранному сопоставлению."""


class ExcelService:
    """Читает и формирует XLSX без исполнения формул книги."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def preview(
        self,
        entity_code: str,
        content: bytes,
        mapping: ExcelMapping | None,
        preview_rows: int,
    ) -> ExcelPreviewRead:
        schema = await self._schema(entity_code)
        formula_workbook, value_workbook = self._workbooks(content)
        sheet = self._sheet(formula_workbook, mapping.sheet_name if mapping else None)
        value_sheet = self._sheet(value_workbook, sheet.title)
        header_row = mapping.header_row if mapping else 1
        headers = self._headers(sheet, header_row)
        technical_headers = self._headers(sheet, 1)
        exported_roundtrip = technical_headers[:4] == [
            "__objectId",
            "__revision",
            "__schemaVersionId",
            "__action",
        ]
        if exported_roundtrip:
            header_row = 2
            headers = self._headers(sheet, header_row)
            fields_by_id = {field.id.hex: field.code for field in schema.fields}
            proposed = {
                headers[index]: fields_by_id[technical.strip().replace("-", "").casefold()]
                for index, technical in enumerate(technical_headers)
                if index < len(headers)
                and technical.strip().replace("-", "").casefold() in fields_by_id
            }
        else:
            proposed = mapping.columns if mapping else self._propose_mapping(headers, schema.fields)
        rows: list[ExcelPreviewRow] = []
        max_row = sheet.max_row or 0
        for row_number in range(header_row + 1, min(max_row, header_row + preview_rows) + 1):
            values, errors = self._row_values(sheet, value_sheet, row_number, headers, proposed)
            parent_object_id = None
            geometry = None
            if mapping is not None:
                parent_object_id, geometry, system_errors = self._row_system_values(
                    sheet,
                    value_sheet,
                    row_number,
                    headers,
                    mapping,
                )
                errors.extend(system_errors)
            if not errors:
                errors.extend(
                    await self._preview_validation_errors(
                        schema,
                        values,
                        parent_object_id,
                        geometry,
                        sheet_name=sheet.title,
                        row_number=row_number,
                        mapping=proposed,
                    )
                )
            rows.append(
                ExcelPreviewRow(
                    row=row_number,
                    values=values,
                    parent_object_id=parent_object_id,
                    geometry=geometry,
                    errors=errors,
                )
            )
        return ExcelPreviewRead(
            sheets=formula_workbook.sheetnames,
            selected_sheet=sheet.title,
            header_row=header_row,
            headers=headers,
            proposed_mapping=proposed,
            total_rows=max(0, max_row - header_row),
            rows=rows,
        )

    async def _preview_validation_errors(
        self,
        schema: EntitySchemaModel,
        values: dict[str, object],
        parent_object_id: UUID | None,
        geometry: GeoJsonGeometry | None,
        *,
        sheet_name: str,
        row_number: int,
        mapping: dict[str, str],
    ) -> list[ExcelCellIssue]:
        """Проверить типы, обязательность, геометрию и родителя строки preview."""

        runtime = RuntimeObjectService(self._session)
        issues: list[dict[str, str | None]] = []
        try:
            await runtime._resolve_parent_object_id(  # noqa: SLF001
                schema,
                parent_object_id,
            )
            normalized = await runtime._normalize_values(  # noqa: SLF001
                schema,
                dict(values),
            )
            issues.extend(
                await runtime._validate(schema, normalized, geometry)  # noqa: SLF001
            )
        except RuntimeValidationError as error:
            issues.extend(error.issues)

        header_by_field = {field_code: header for header, field_code in mapping.items()}
        return [
            ExcelCellIssue(
                sheet=sheet_name,
                row=row_number,
                column=header_by_field.get(
                    str(issue.get("fieldCode")),
                    "Системные поля",
                ),
                field_code=issue.get("fieldCode"),
                code=str(issue.get("code") or "invalid_value"),
                message=str(issue.get("message") or "Некорректное значение"),
            )
            for issue in issues
        ]

    async def to_objects(
        self, entity_code: str, content: bytes, mapping: ExcelMapping
    ) -> list[EntityObjectCreate]:
        schema = await self._schema(entity_code)
        formula_workbook, value_workbook = self._workbooks(content)
        sheet = self._sheet(formula_workbook, mapping.sheet_name)
        value_sheet = self._sheet(value_workbook, sheet.title)
        headers = self._headers(sheet, mapping.header_row)
        allowed_codes = {field.code for field in schema.fields if not field.read_only and not field.archived}
        invalid_codes = sorted(set(mapping.columns.values()) - allowed_codes)
        if invalid_codes:
            raise ExcelValidationError(
                "Сопоставление содержит недоступные поля: " + ", ".join(invalid_codes)
            )
        objects: list[EntityObjectCreate] = []
        issues: list[ExcelCellIssue] = []
        for row_number in range(mapping.header_row + 1, (sheet.max_row or 0) + 1):
            values, row_issues = self._row_values(
                sheet, value_sheet, row_number, headers, mapping.columns
            )
            parent_object_id, geometry, system_issues = self._row_system_values(
                sheet,
                value_sheet,
                row_number,
                headers,
                mapping,
            )
            row_issues.extend(system_issues)
            issues.extend(row_issues)
            if not values and geometry is None and parent_object_id is None and not row_issues:
                continue
            objects.append(
                EntityObjectCreate(
                    values=values,
                    parent_object_id=parent_object_id,
                    geometry=geometry,
                )
            )
        if issues:
            first = issues[0]
            raise ExcelValidationError(
                f"{first.sheet}, строка {first.row}, столбец {first.column}: {first.message}"
            )
        if not objects:
            raise ExcelValidationError("В выбранном листе нет строк данных")
        return objects

    async def to_change_set_items(
        self, entity_code: str, content: bytes, mapping: ExcelMapping
    ) -> list[ChangeSetItemCreate]:
        """Преобразовать произвольный или ранее экспортированный XLSX в ChangeSet."""

        schema = await self._schema(entity_code)
        current_schema_version_id = await self._schema_version_id(schema.id)
        formula_workbook, value_workbook = self._workbooks(content)
        sheet = self._sheet(formula_workbook, mapping.sheet_name)
        value_sheet = self._sheet(value_workbook, sheet.title)
        technical_headers = self._headers(sheet, 1)
        if technical_headers[:4] != [
            "__objectId",
            "__revision",
            "__schemaVersionId",
            "__action",
        ]:
            if mapping.system.object_id_column:
                headers = self._headers(sheet, mapping.header_row)
                if mapping.system.object_id_column not in headers:
                    raise ExcelValidationError("Колонка objectId не найдена")
                if (
                    mapping.system.action_column
                    and mapping.system.action_column not in headers
                ):
                    raise ExcelValidationError("Колонка действия не найдена")
                object_index = headers.index(mapping.system.object_id_column)
                action_index = (
                    headers.index(mapping.system.action_column)
                    if mapping.system.action_column
                    else None
                )
                parsed_rows: list[
                    tuple[UUID | None, str, dict[str, object], UUID | None, object | None]
                ] = []
                for row_number in range(mapping.header_row + 1, (sheet.max_row or 0) + 1):
                    values, errors = self._row_values(
                        sheet,
                        value_sheet,
                        row_number,
                        headers,
                        mapping.columns,
                    )
                    parent_id, geometry, system_errors = self._row_system_values(
                        sheet,
                        value_sheet,
                        row_number,
                        headers,
                        mapping,
                    )
                    errors.extend(system_errors)
                    if errors:
                        first = errors[0]
                        raise ExcelValidationError(
                            f"{first.sheet}, строка {first.row}: {first.message}"
                        )
                    raw_id = self._import_cell_value(
                        sheet[row_number][object_index],
                        value_sheet[row_number][object_index],
                        sheet.title,
                        row_number,
                    )
                    try:
                        object_id = UUID(str(raw_id).strip()) if raw_id not in (None, "") else None
                    except ValueError as error:
                        raise ExcelValidationError(
                            f"{sheet.title}, строка {row_number}: objectId должен быть UUID"
                        ) from error
                    raw_action = (
                        self._import_cell_value(
                            sheet[row_number][action_index],
                            value_sheet[row_number][action_index],
                            sheet.title,
                            row_number,
                        )
                        if action_index is not None
                        else None
                    )
                    operation = str(raw_action or ("update" if object_id else "create")).strip().casefold()
                    aliases = {"создать": "create", "изменить": "update", "архивировать": "archive"}
                    operation = aliases.get(operation, operation)
                    if operation not in {"create", "update", "archive"}:
                        raise ExcelValidationError(
                            f"{sheet.title}, строка {row_number}: неизвестное действие {operation}"
                        )
                    parsed_rows.append((object_id, operation, values, parent_id, geometry))
                object_ids = {row[0] for row in parsed_rows if row[0] is not None}
                current_by_id = {
                    item.id: item
                    for item in await self._session.scalars(
                        select(EntityObjectModel).where(
                            EntityObjectModel.entity_schema_id == schema.id,
                            EntityObjectModel.id.in_(object_ids),
                        )
                    )
                }
                missing = object_ids - set(current_by_id)
                if missing:
                    raise ExcelValidationError(
                        "Объекты из objectId не найдены: "
                        + ", ".join(str(item) for item in sorted(missing, key=str))
                    )
                return [
                    ChangeSetItemCreate(
                        object_id=object_id,
                        operation=operation,
                        base_revision=(
                            current_by_id[object_id].revision if object_id is not None else None
                        ),
                        values={} if operation == "archive" else values,
                        parent_object_id=parent_id,
                        geometry=geometry,
                    )
                    for object_id, operation, values, parent_id, geometry in parsed_rows
                ]
            objects = await self.to_objects(entity_code, content, mapping)
            return [
                ChangeSetItemCreate(
                    operation="create",
                    values=item.values,
                    parent_object_id=item.parent_object_id,
                    geometry=item.geometry,
                )
                for item in objects
            ]

        field_by_id = {
            field.id.hex: field
            for field in schema.fields
            if not field.archived
        }
        column_fields = {
            index: field_by_id[header.strip().replace("-", "").casefold()]
            for index, header in enumerate(technical_headers[4:], start=4)
            if header.strip().replace("-", "").casefold() in field_by_id
        }
        items: list[ChangeSetItemCreate] = []
        for row_number in range(3, (sheet.max_row or 0) + 1):
            formula_row = sheet[row_number]
            value_row = value_sheet[row_number]
            if not any(cell.value is not None for cell in formula_row):
                continue
            raw_object_id = formula_row[0].value
            raw_revision = formula_row[1].value
            raw_schema_version_id = formula_row[2].value
            raw_action = str(formula_row[3].value or "").strip().casefold()
            values: dict[str, object] = {}
            for index, field in column_fields.items():
                if field.read_only:
                    continue
                value = self._import_cell_value(
                    formula_row[index], value_row[index], sheet.title, row_number
                )
                values[field.code] = value
            if raw_action in {"archive", "архив", "архивировать"}:
                if raw_object_id is None:
                    raise ExcelValidationError(
                        f"{sheet.title}, строка {row_number}: для архивирования нужен __objectId"
                    )
                items.append(
                    ChangeSetItemCreate(
                        object_id=UUID(str(raw_object_id)),
                        operation="archive",
                        base_revision=int(raw_revision) if raw_revision is not None else None,
                    )
                )
                continue
            if raw_object_id is None or str(raw_object_id).strip() == "":
                items.append(ChangeSetItemCreate(operation="create", values=values))
                continue
            if raw_revision is None or raw_schema_version_id is None:
                raise ExcelValidationError(
                    f"{sheet.title}, строка {row_number}: "
                    "отсутствует __revision или __schemaVersionId"
                )
            try:
                object_id = UUID(str(raw_object_id))
                revision = int(raw_revision)
                schema_version_id = UUID(str(raw_schema_version_id))
            except (TypeError, ValueError) as error:
                raise ExcelValidationError(
                    f"{sheet.title}, строка {row_number}: "
                    "некорректный ID, ревизия или версия схемы"
                ) from error
            if schema_version_id != current_schema_version_id:
                raise ExcelValidationError(
                    f"{sheet.title}, строка {row_number}: структура сущности изменилась; "
                    "выгрузите новый XLSX и повторите изменения"
                )
            items.append(
                ChangeSetItemCreate(
                    object_id=object_id,
                    operation="update",
                    base_revision=revision,
                    values=values,
                )
            )
        if not items:
            raise ExcelValidationError("В выбранном листе нет строк данных")
        return items

    async def export(
        self,
        entity_code: str,
        *,
        actor: ActorContext,
    ) -> tuple[str, bytes]:
        schema = await self._schema(entity_code)
        conditions = [
            EntityObjectModel.entity_schema_id == schema.id,
            EntityObjectModel.status != "archived",
            await AuthorizationService(self._session).readable_object_condition(
                actor,
                schema=schema,
            ),
        ]
        total = int(
            await self._session.scalar(
                select(func.count()).select_from(EntityObjectModel).where(*conditions)
            )
            or 0
        )
        if total > settings.excel_max_rows:
            raise ExcelValidationError(
                f"В выборке {total} объектов, лимит одной XLSX-выгрузки — "
                f"{settings.excel_max_rows}. Сузьте выборку или используйте экспорт выбранных данных."
            )
        objects = (
            await self._session.scalars(
                select(EntityObjectModel)
                .where(*conditions)
                .order_by(EntityObjectModel.created_at)
            )
        ).all()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Объекты"
        fields = [
            field
            for field in schema.fields
            if not field.archived and self._field_is_readable(field, actor.roles)
        ]
        sheet.append([field.name for field in fields])
        runtime = RuntimeObjectService(self._session)
        for item in objects:
            display_values = await runtime._display_values(  # noqa: SLF001
                schema,
                item.values,
            )
            sheet.append(
                [
                    *[
                        self._safe_excel_value(
                            display_values.get(field.code, item.values.get(field.code))
                        )
                        for field in fields
                    ],
                ]
            )
        self._format_export_sheet(sheet)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return f"{schema.code}.xlsx", buffer.getvalue()

    async def export_tree(
        self,
        entity_code: str,
        selection: RegistryTreeSearch | ExcelExportRequest,
        *,
        actor: ActorContext,
        max_rows: int | None = None,
    ) -> tuple[str, bytes, int]:
        """Выгрузить выбранные колонки реестра и подреестров в одну книгу."""

        row_limit = max_rows or settings.excel_max_rows
        root_schema = await self._schema(entity_code)
        tree = await RuntimeObjectService(self._session).list_tree(
            entity_code,
            selection,
            actor=actor,
            max_returned_rows=row_limit,
        )
        root_fields = self._selected_fields(
            root_schema,
            selection.columns,
            actor_roles=actor.roles,
        )
        workbook = Workbook()
        root_sheet = workbook.active
        root_sheet.title = self._sheet_title(root_schema.name, "Реестр")
        root_sheet.append([field.name for field in root_fields])
        exported_rows = len(tree.data)
        if exported_rows > row_limit:
            raise ExcelValidationError(
                f"В выгрузке {exported_rows} строк, предел фоновой "
                f"XLSX-выгрузки — {row_limit}"
            )
        for item in tree.data:
            root_sheet.append(
                [
                    *[
                        self._safe_excel_value(
                            item.display_values.get(field.code, item.values.get(field.code))
                        )
                        for field in root_fields
                    ],
                ]
            )
        self._format_export_sheet(root_sheet)

        for child_selection in selection.children:
            child_schema = await self._schema(child_selection.entity_code)
            child_fields = self._selected_fields(
                child_schema,
                child_selection.columns,
                actor_roles=actor.roles,
            )
            sheet = workbook.create_sheet(
                self._unique_sheet_title(workbook, child_schema.name)
            )
            sheet.append(
                [
                    *[f"{root_schema.name}: {field.name}" for field in root_fields],
                    *[field.name for field in child_fields],
                ]
            )
            for parent in tree.data:
                parent_values = [
                    self._safe_excel_value(
                        parent.display_values.get(field.code, parent.values.get(field.code))
                    )
                    for field in root_fields
                ]
                for child in parent.children.get(child_schema.code, []):
                    exported_rows += 1
                    if exported_rows > row_limit:
                        raise ExcelValidationError(
                            f"Суммарная выгрузка реестра и подреестров "
                            f"превышает лимит {row_limit} строк"
                        )
                    sheet.append(
                        [
                            *parent_values,
                            *[
                                self._safe_excel_value(
                                    child.display_values.get(
                                        field.code,
                                        child.values.get(field.code),
                                    )
                                )
                                for field in child_fields
                            ],
                        ]
                    )
            self._format_export_sheet(sheet)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return (
            f"{root_schema.code}-with-children.xlsx",
            buffer.getvalue(),
            exported_rows,
        )

    @staticmethod
    def _selected_fields(
        schema: EntitySchemaModel,
        columns: list[str],
        *,
        actor_roles: frozenset[str] | None = None,
    ) -> list[EntityFieldModel]:
        available = {
            field.code: field
            for field in schema.fields
            if not field.archived
            and ExcelService._field_is_readable(field, actor_roles)
        }
        codes = columns or [
            field.code
            for field in schema.fields
            if field.list_visible
            and not field.archived
            and ExcelService._field_is_readable(field, actor_roles)
        ]
        unknown = sorted(set(codes) - set(available))
        if unknown:
            raise ExcelValidationError(
                "Неизвестные или архивные колонки: " + ", ".join(unknown)
            )
        return [available[code] for code in codes]

    @staticmethod
    def _field_is_readable(
        field: EntityFieldModel,
        actor_roles: frozenset[str] | None,
    ) -> bool:
        roles = (field.access_rules or {}).get("read")
        return (
            roles is None
            or bool(actor_roles and set(roles).intersection(actor_roles))
        )

    @staticmethod
    def _sheet_title(value: str, fallback: str) -> str:
        title = "".join("_" if char in "[]:*?/\\" else char for char in value).strip()
        return (title or fallback)[:31]

    @classmethod
    def _unique_sheet_title(cls, workbook: Workbook, value: str) -> str:
        base = cls._sheet_title(value, "Подреестр")
        candidate = base
        suffix = 2
        while candidate in workbook.sheetnames:
            marker = f" ({suffix})"
            candidate = f"{base[: 31 - len(marker)]}{marker}"
            suffix += 1
        return candidate

    @staticmethod
    def _format_export_sheet(sheet: Worksheet) -> None:
        """Сделать пользовательскую выгрузку читаемой сразу после открытия."""

        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Arial", size=10, color="1F2937")
        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        body_alignment = Alignment(vertical="center", wrap_text=True)

        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.row_dimensions[1].height = 34

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 24
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    # Значение остаётся строкой без добавления видимого апострофа.
                    cell.data_type = "s"
                cell.font = body_font
                cell.alignment = body_alignment
                if isinstance(cell.value, int | float):
                    cell.number_format = "#,##0.##"

        for column_index in range(1, sheet.max_column + 1):
            values = [
                str(sheet.cell(row=row_index, column=column_index).value or "")
                for row_index in range(1, sheet.max_row + 1)
            ]
            content_width = max((len(value) for value in values), default=0) + 2
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(content_width, 14),
                44,
            )

    def objects_json(self, objects: list[EntityObjectCreate]) -> bytes:
        return orjson.dumps([item.model_dump(mode="json", by_alias=True) for item in objects])

    def _workbooks(self, content: bytes):
        if not content:
            raise ExcelValidationError("XLSX-файл пуст")
        if len(content) > settings.excel_max_file_size_bytes:
            raise ExcelValidationError("XLSX-файл превышает допустимый размер")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
                required_parts = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
                if not required_parts.issubset(names):
                    raise ExcelValidationError(
                        "Файл не содержит обязательные части книги XLSX"
                    )
                if sum(item.file_size for item in archive.infolist()) > settings.excel_max_uncompressed_bytes:
                    raise ExcelValidationError("Распакованный XLSX превышает допустимый размер")
            formula_workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            value_workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_links=False
            )
        except ExcelValidationError:
            raise
        except (
            zipfile.BadZipFile,
            InvalidFileException,
            KeyError,
            OSError,
            ParseError,
            ValueError,
        ) as error:
            raise ExcelValidationError("Файл не является корректной книгой XLSX") from error
        if len(formula_workbook.sheetnames) > settings.excel_max_sheets:
            raise ExcelValidationError("В XLSX слишком много листов")
        for sheet in formula_workbook.worksheets:
            max_row = sheet.max_row or 0
            max_column = sheet.max_column or 0
            if max_row > settings.excel_max_rows + 2:
                raise ExcelValidationError(f"Лист «{sheet.title}» превышает лимит строк")
            if max_column > settings.excel_max_columns:
                raise ExcelValidationError(f"Лист «{sheet.title}» превышает лимит столбцов")
        return formula_workbook, value_workbook

    @staticmethod
    def _sheet(workbook, sheet_name: str | None) -> Worksheet:
        if sheet_name is None:
            return workbook[workbook.sheetnames[0]]
        if sheet_name not in workbook.sheetnames:
            raise ExcelValidationError(f"Лист «{sheet_name}» не найден")
        return workbook[sheet_name]

    @staticmethod
    def _headers(sheet: Worksheet, header_row: int) -> list[str]:
        if header_row > (sheet.max_row or 0):
            raise ExcelValidationError("Строка заголовков находится за пределами листа")
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[header_row]]
        if not any(headers):
            raise ExcelValidationError("Строка заголовков пуста")
        return headers

    @staticmethod
    def _propose_mapping(headers: list[str], fields: list[EntityFieldModel]) -> dict[str, str]:
        by_token: dict[str, str] = {}
        for field in fields:
            if field.archived or field.read_only:
                continue
            by_token[field.code.casefold()] = field.code
            by_token[field.name.strip().casefold()] = field.code
            by_token[field.id.hex.casefold()] = field.code
        return {
            header: by_token[token]
            for header in headers
            if header and (token := header.strip().casefold()) in by_token
        }

    @staticmethod
    def _row_values(
        formula_sheet: Worksheet,
        value_sheet: Worksheet,
        row_number: int,
        headers: list[str],
        mapping: dict[str, str],
    ) -> tuple[dict[str, object], list[ExcelCellIssue]]:
        values: dict[str, object] = {}
        issues: list[ExcelCellIssue] = []
        formula_row = formula_sheet[row_number]
        value_row = value_sheet[row_number]
        for index, header in enumerate(headers):
            field_code = mapping.get(header)
            if not field_code or index >= len(formula_row):
                continue
            cell: Cell = formula_row[index]
            value_cell: Cell = value_row[index]
            if cell.data_type == "f":
                if value_cell.value is None:
                    issues.append(
                        ExcelCellIssue(
                            sheet=formula_sheet.title,
                            row=row_number,
                            column=cell.coordinate,
                            field_code=field_code,
                            code="formula_result_missing",
                            message=(
                                "В файле нет сохранённого результата формулы; "
                                "пересчитайте и сохраните книгу в Excel или LibreOffice"
                            ),
                        )
                    )
                    continue
                value = value_cell.value
            else:
                value = cell.value
            if value is None:
                continue
            if isinstance(value, datetime | date):
                value = value.isoformat()
            elif isinstance(value, Decimal):
                value = str(value)
            values[field_code] = value
        return values, issues

    @classmethod
    def _row_system_values(
        cls,
        formula_sheet: Worksheet,
        value_sheet: Worksheet,
        row_number: int,
        headers: list[str],
        mapping: ExcelMapping,
    ) -> tuple[UUID | None, object | None, list[ExcelCellIssue]]:
        system = mapping.system
        issues: list[ExcelCellIssue] = []
        parent_object_id = system.fixed_parent_object_id

        def read(header: str | None, code: str) -> object | None:
            if not header:
                return None
            if header not in headers:
                issues.append(
                    ExcelCellIssue(
                        sheet=formula_sheet.title,
                        row=row_number,
                        column=header,
                        code="system_column_not_found",
                        message=f"Системный столбец «{header}» не найден",
                    )
                )
                return None
            index = headers.index(header)
            try:
                return cls._import_cell_value(
                    formula_sheet[row_number][index],
                    value_sheet[row_number][index],
                    formula_sheet.title,
                    row_number,
                )
            except ExcelValidationError as error:
                issues.append(
                    ExcelCellIssue(
                        sheet=formula_sheet.title,
                        row=row_number,
                        column=header,
                        code=code,
                        message=str(error),
                    )
                )
                return None

        if system.parent_object_column:
            raw_parent = read(system.parent_object_column, "invalid_parent")
            if raw_parent not in (None, ""):
                try:
                    parent_object_id = UUID(str(raw_parent).strip())
                except ValueError:
                    issues.append(
                        ExcelCellIssue(
                            sheet=formula_sheet.title,
                            row=row_number,
                            column=system.parent_object_column,
                            code="parent_not_found",
                            message="ID родительского объекта должен быть UUID",
                        )
                    )

        geometry = None
        if system.geometry_column:
            raw_geometry = read(system.geometry_column, "invalid_geometry")
            if raw_geometry not in (None, ""):
                try:
                    payload = (
                        json.loads(raw_geometry)
                        if isinstance(raw_geometry, str)
                        else raw_geometry
                    )
                    from app.modules.objects.api.schemas import GeoJsonGeometry

                    geometry = GeoJsonGeometry.model_validate(payload)
                except (json.JSONDecodeError, ValidationError, TypeError, ValueError):
                    issues.append(
                        ExcelCellIssue(
                            sheet=formula_sheet.title,
                            row=row_number,
                            column=system.geometry_column,
                            code="invalid_geometry",
                            message="Геометрия должна быть корректным GeoJSON",
                        )
                    )
        elif system.latitude_column and system.longitude_column:
            raw_latitude = read(system.latitude_column, "invalid_coordinates")
            raw_longitude = read(system.longitude_column, "invalid_coordinates")
            if raw_latitude not in (None, "") or raw_longitude not in (None, ""):
                try:
                    latitude = float(raw_latitude)
                    longitude = float(raw_longitude)
                    if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
                        raise ValueError
                    from app.modules.objects.api.schemas import GeoJsonGeometry

                    geometry = GeoJsonGeometry(
                        type="Point",
                        coordinates=[longitude, latitude],
                    )
                except (TypeError, ValueError):
                    issues.append(
                        ExcelCellIssue(
                            sheet=formula_sheet.title,
                            row=row_number,
                            column=f"{system.longitude_column}/{system.latitude_column}",
                            code="invalid_coordinates",
                            message="Ожидаются долгота [-180; 180] и широта [-90; 90]",
                        )
                    )
        return parent_object_id, geometry, issues

    @staticmethod
    def _safe_excel_value(value: object) -> object:
        if isinstance(value, list | dict):
            return json.dumps(value, ensure_ascii=False)
        return value

    @staticmethod
    def _import_cell_value(
        formula_cell: Cell,
        value_cell: Cell,
        sheet_name: str,
        row_number: int,
    ) -> object:
        if formula_cell.data_type == "f":
            if value_cell.value is None:
                raise ExcelValidationError(
                    f"{sheet_name}, строка {row_number}, столбец {formula_cell.coordinate}: "
                    "нет сохранённого результата формулы"
                )
            value = value_cell.value
        else:
            value = formula_cell.value
        if isinstance(value, datetime | date):
            return value.isoformat()
        if isinstance(value, Decimal):
            return format(value, "f")
        return value

    async def _schema(self, entity_code: str) -> EntitySchemaModel:
        schema = await self._session.scalar(
            select(EntitySchemaModel).where(
                EntitySchemaModel.code == entity_code,
                EntitySchemaModel.status == "active",
            )
        )
        if schema is None:
            raise RuntimeEntityNotFound
        return schema

    async def _schema_version_id(self, entity_id: UUID) -> UUID:
        version_id = await self._session.scalar(
            select(EntitySchemaVersionModel.id)
            .where(EntitySchemaVersionModel.entity_schema_id == entity_id)
            .order_by(EntitySchemaVersionModel.version.desc())
            .limit(1)
        )
        if version_id is None:
            raise ExcelValidationError("У сущности отсутствует опубликованная версия схемы")
        return version_id


class ExcelImportPlanService:
    """Готовит проверяемый Excel-план и применяет его через атомарный ChangeSet."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def create(
        self,
        entity_code: str,
        source_name: str,
        content: bytes,
        mapping: ExcelMapping,
        actor: ActorContext,
        *,
        idempotency_key: str | None,
    ) -> ExcelImportPlanRead:
        excel = ExcelService(self._session)
        items = await excel.to_change_set_items(entity_code, content, mapping)
        await self._authorize_items(entity_code, items, actor)
        await self._session.rollback()
        change_set = await ChangeSetService(self._session).create(
            ChangeSetCreate(
                entity_code=entity_code,
                source="excel",
                idempotency_key=idempotency_key,
                items=items,
                metadata={"sourceName": source_name, "mode": "preview_plan"},
            ),
            actor,
        )
        current_ids = [item.object_id for item in change_set.items if item.object_id]
        current_by_id: dict[UUID, EntityObjectModel] = {}
        if current_ids:
            current_by_id = {
                item.id: item
                for item in await self._session.scalars(
                    select(EntityObjectModel).where(EntityObjectModel.id.in_(current_ids))
                )
            }
        rows: list[dict[str, object]] = []
        summary = {
            "new": 0,
            "changed": 0,
            "unchanged": 0,
            "archived": 0,
            "errors": 0,
            "conflicts": 0,
        }
        for index, item in enumerate(change_set.items):
            current = current_by_id.get(item.object_id) if item.object_id else None
            base_values = dict(current.values) if current else None
            if item.validation_errors:
                state = "error"
                summary["errors"] += 1
            elif item.operation == "create":
                state = "new"
                summary["new"] += 1
            elif item.operation == "archive":
                state = "archived"
                summary["archived"] += 1
            elif base_values == item.proposed_values:
                state = "unchanged"
                summary["unchanged"] += 1
            else:
                state = "changed"
                summary["changed"] += 1
            fields = []
            for code in sorted(set(base_values or {}) | set(item.proposed_values)):
                before = (base_values or {}).get(code)
                after = item.proposed_values.get(code)
                if before != after:
                    fields.append(
                        {"fieldCode": code, "base": before, "file": after, "current": before}
                    )
            rows.append(
                {
                    "rowIndex": index,
                    "objectId": str(item.object_id) if item.object_id else None,
                    "operation": item.operation,
                    "state": state,
                    "baseRevision": item.base_revision,
                    "base": base_values,
                    "file": item.proposed_values,
                    "current": base_values,
                    "fields": fields,
                    "validationErrors": [issue.model_dump(mode="json", by_alias=True) for issue in item.validation_errors],
                }
            )
        schema = await excel._schema(entity_code)  # noqa: SLF001
        await self._session.rollback()
        async with self._session.begin():
            model = ExcelImportPlanModel(
                entity_schema_id=schema.id,
                schema_version_id=change_set.schema_version_id,
                created_by=actor.id,
                status="ready" if change_set.status == "validated" else "draft",
                source_name=source_name,
                mapping=mapping.model_dump(mode="json", by_alias=True),
                summary=summary,
                rows=rows,
                decisions={"changeSetId": str(change_set.id)},
                expires_at=datetime.now(UTC) + timedelta(hours=24),
            )
            self._session.add(model)
            await self._session.flush()
            await self._session.refresh(model)
        return self._response(model, limit=100, offset=0)

    async def _authorize_items(
        self,
        entity_code: str,
        items: list[ChangeSetItemCreate],
        actor: ActorContext,
    ) -> None:
        """Проверить предметные операции плана до сохранения его снимка."""

        authorization = AuthorizationService(self._session)
        checked: set[tuple[str, UUID | None]] = set()
        for item in items:
            action = {
                "create": "create",
                "update": "update",
                "archive": "archive",
                "confirm": "confirm",
            }[item.operation]
            key = (action, item.object_id)
            if key not in checked:
                await authorization.require_entity_code(
                    actor,
                    action,
                    entity_code,
                    object_id=item.object_id,
                )
                checked.add(key)

        parent_ids = {
            item.parent_object_id
            for item in items
            if item.parent_object_id is not None
        }
        if not parent_ids:
            return
        schema = await self._session.scalar(
            select(EntitySchemaModel).where(EntitySchemaModel.code == entity_code)
        )
        if schema is None or schema.parent_entity_schema_id is None:
            raise ExcelValidationError("Для корневой сущности нельзя указать родительский объект")
        parent_code = await self._session.scalar(
            select(EntitySchemaModel.code).where(
                EntitySchemaModel.id == schema.parent_entity_schema_id,
                EntitySchemaModel.status == "active",
            )
        )
        if parent_code is None:
            raise ExcelValidationError("Родительская сущность не опубликована")
        for parent_id in parent_ids:
            await authorization.require_entity_code(
                actor,
                "read",
                parent_code,
                object_id=parent_id,
            )

    async def get(
        self,
        plan_id: UUID,
        actor_id: UUID,
        *,
        limit: int,
        offset: int,
    ) -> ExcelImportPlanRead:
        model = await self._session.scalar(
            select(ExcelImportPlanModel).where(
                ExcelImportPlanModel.id == plan_id,
                ExcelImportPlanModel.created_by == actor_id,
            )
        )
        if model is None:
            raise ExcelValidationError("План импорта не найден")
        if model.expires_at <= datetime.now(UTC) and model.status not in {"applied", "rejected"}:
            model.status = "expired"
            await self._session.commit()
        return self._response(model, limit=limit, offset=offset)

    async def decide(
        self,
        plan_id: UUID,
        payload: ExcelImportPlanDecision,
        actor: ActorContext,
    ) -> ExcelImportPlanRead:
        model = await self._session.scalar(
            select(ExcelImportPlanModel).where(
                ExcelImportPlanModel.id == plan_id,
                ExcelImportPlanModel.created_by == actor.id,
            )
        )
        if model is None:
            raise ExcelValidationError("План импорта не найден")
        if model.expires_at <= datetime.now(UTC):
            model.status = "expired"
            await self._session.commit()
            raise ExcelValidationError("Срок действия плана импорта истёк")
        if model.status == "applied":
            return self._response(model, limit=100, offset=0)
        if payload.action in {"reject", "cancel"}:
            model.status = "rejected"
            model.decisions = {**model.decisions, "action": payload.action}
            await self._session.commit()
            return self._response(model, limit=100, offset=0)
        if model.status != "ready":
            raise ExcelValidationError("План содержит ошибки и не может быть применён")
        change_set_id = UUID(str(model.decisions["changeSetId"]))
        await self._authorize_change_set(change_set_id, actor)
        await self._session.rollback()
        await self._prepare_resolutions(
            change_set_id,
            payload.resolutions,
        )
        try:
            await ChangeSetService(self._session).apply(change_set_id, actor)
        except ChangeSetConflict as error:
            raise ExcelValidationError(str(error)) from error
        except ChangeSetValidationError as error:
            raise ExcelValidationError(str(error)) from error
        async with self._session.begin():
            model = await self._session.scalar(
                select(ExcelImportPlanModel)
                .where(ExcelImportPlanModel.id == plan_id)
                .with_for_update()
            )
            model.status = "applied"
            model.decisions = {
                **model.decisions,
                "action": "apply",
                "resolutions": payload.resolutions,
                "idempotencyKey": payload.idempotency_key,
            }
        return self._response(model, limit=100, offset=0)

    async def _authorize_change_set(
        self,
        change_set_id: UUID,
        actor: ActorContext,
    ) -> None:
        row = (
            await self._session.execute(
                select(ChangeSetModel, EntitySchemaModel.code)
                .join(
                    EntitySchemaModel,
                    EntitySchemaModel.id == ChangeSetModel.entity_schema_id,
                )
                .where(ChangeSetModel.id == change_set_id)
            )
        ).one_or_none()
        if row is None:
            raise ExcelValidationError("Проверенный набор изменений недоступен")
        change_set, entity_code = row
        items = list(
            await self._session.scalars(
                select(ChangeSetItemModel).where(
                    ChangeSetItemModel.change_set_id == change_set_id
                )
            )
        )
        authorization = AuthorizationService(self._session)
        checked: set[tuple[str, UUID | None]] = set()
        for item in items:
            action = {
                "create": "create",
                "update": "update",
                "archive": "archive",
                "confirm": "confirm",
            }[item.operation]
            key = (action, item.object_id)
            if key in checked:
                continue
            await authorization.require_entity_code(
                actor,
                action,
                entity_code,
                object_id=item.object_id,
            )
            checked.add(key)

        parent_ids = {
            item.parent_object_id
            for item in items
            if item.parent_object_id is not None
        }
        if not parent_ids:
            return
        schema = await self._session.get(
            EntitySchemaModel,
            change_set.entity_schema_id,
        )
        parent_code = (
            await self._session.scalar(
                select(EntitySchemaModel.code).where(
                    EntitySchemaModel.id == schema.parent_entity_schema_id,
                    EntitySchemaModel.status == "active",
                )
            )
            if schema is not None and schema.parent_entity_schema_id is not None
            else None
        )
        if parent_code is None:
            raise ExcelValidationError("Родительская сущность не опубликована")
        for parent_id in parent_ids:
            await authorization.require_entity_code(
                actor,
                "read",
                parent_code,
                object_id=parent_id,
            )

    async def _prepare_resolutions(
        self,
        change_set_id: UUID,
        resolutions: dict[str, str],
    ) -> None:
        async with self._session.begin():
            change_set = await self._session.scalar(
                select(ChangeSetModel)
                .where(ChangeSetModel.id == change_set_id)
                .with_for_update()
            )
            if change_set is None or change_set.status != "validated":
                raise ExcelValidationError("Проверенный набор изменений недоступен")
            items = list(
                await self._session.scalars(
                    select(ChangeSetItemModel)
                    .where(ChangeSetItemModel.change_set_id == change_set_id)
                    .order_by(ChangeSetItemModel.id)
                    .with_for_update()
                )
            )
            for row_index, item in enumerate(items):
                if item.operation not in {"update", "archive"} or item.object_id is None:
                    continue
                current = await self._session.scalar(
                    select(EntityObjectModel)
                    .where(EntityObjectModel.id == item.object_id)
                    .with_for_update()
                )
                if current is None:
                    raise ExcelValidationError(
                        f"Объект строки {row_index} был удалён после проверки"
                    )
                if item.base_revision == current.revision:
                    continue
                if item.operation == "archive":
                    key = f"{row_index}.__row"
                    if resolutions.get(key) != "file":
                        raise ExcelValidationError(
                            f"Строка {row_index} изменилась; передайте resolution {key}=file"
                        )
                    item.base_revision = current.revision
                    continue
                proposed = dict(item.proposed_values)
                for code in sorted(set(proposed) | set(current.values)):
                    if proposed.get(code) == current.values.get(code):
                        continue
                    key = f"{row_index}.{code}"
                    resolution = resolutions.get(key)
                    if resolution is None:
                        raise ExcelValidationError(
                            f"Поле {code} строки {row_index} изменилось; требуется resolution"
                        )
                    if resolution == "current":
                        proposed[code] = current.values.get(code)
                item.proposed_values = proposed
                item.base_revision = current.revision

    @staticmethod
    def _response(
        model: ExcelImportPlanModel,
        *,
        limit: int,
        offset: int,
    ) -> ExcelImportPlanRead:
        return ExcelImportPlanRead(
            id=model.id,
            entity_id=model.entity_schema_id,
            schema_version_id=model.schema_version_id,
            source_name=model.source_name,
            status=model.status,
            summary=model.summary,
            rows=model.rows[offset : offset + limit],
            total_rows=len(model.rows),
            limit=limit,
            offset=offset,
            expires_at=model.expires_at,
            created_at=model.created_at,
            updated_at=model.updated_at,
        )


class ImportProfileService:
    """Сохраняет подтверждённые пользователем Excel-сопоставления."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def create(self, payload: ImportProfileCreate, owner_id: UUID) -> ImportProfileRead:
        async with self._session.begin():
            model = ImportProfileModel(
                entity_schema_id=payload.entity_schema_id,
                owner_id=owner_id,
                name=payload.name,
                sheet_name=payload.mapping.sheet_name,
                header_row=payload.mapping.header_row,
                mapping=payload.mapping.columns,
            )
            self._session.add(model)
            await self._session.flush()
            await self._session.refresh(model)
        return ImportProfileRead.model_validate(model, from_attributes=True)

    async def list(self, entity_id: UUID, owner_id: UUID) -> list[ImportProfileRead]:
        rows = (
            await self._session.scalars(
                select(ImportProfileModel)
                .where(
                    ImportProfileModel.entity_schema_id == entity_id,
                    ImportProfileModel.owner_id == owner_id,
                )
                .order_by(ImportProfileModel.name)
            )
        ).all()
        return [ImportProfileRead.model_validate(row, from_attributes=True) for row in rows]

    async def delete(self, profile_id: UUID, owner_id: UUID) -> None:
        async with self._session.begin():
            result = await self._session.execute(
                delete(ImportProfileModel).where(
                    ImportProfileModel.id == profile_id,
                    ImportProfileModel.owner_id == owner_id,
                )
            )
            if result.rowcount == 0:
                raise ExcelValidationError("Шаблон импорта не найден")
